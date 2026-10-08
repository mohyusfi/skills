"""
docx_styler.py
Modul perender dokumen Microsoft Word (.docx) dengan styling akademik standar:
- Kertas A4, Margin 4-4-3-3 cm
- Times New Roman 12 pt, 1.5 Line Spacing
- Heading berhierarki native Word (kompatibel penuh Navigation Pane)
- Paragraf Justified berindentasi 1.27 cm
- Poin List rapi
- Placeholder Gambar
- Daftar Pustaka berformat hanging indent (APA 7th)
"""

import os
import re
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

from kbbi_italicizer import parse_runs_for_italics

def set_paragraph_spacing(paragraph, line_spacing=1.5, space_before_pt=0, space_after_pt=0, keep_with_next=False):
    """Mengatur spasi baris 1.5 dan spasi sebelum/sesudah."""
    p_format = paragraph.paragraph_format
    p_format.line_spacing = line_spacing
    p_format.space_before = Pt(space_before_pt)
    p_format.space_after = Pt(space_after_pt)
    if keep_with_next:
        p_format.keep_with_next = True

def set_outline_level(paragraph, level: int):
    """Menyematkan w:outlineLvl pada XML paragraf agar muncul di MS Word Navigation Pane."""
    pPr = paragraph._element.get_or_add_pPr()
    # Hapus outlineLvl lama jika ada
    for child in list(pPr):
        if child.tag.endswith('outlineLvl'):
            pPr.remove(child)
    outline = OxmlElement('w:outlineLvl')
    outline.set(qn('w:val'), str(level))
    pPr.append(outline)

def add_runs(paragraph, text, is_heading=False, base_bold=False, base_italic=False, font_size_pt=12):
    """Menambahkan run ke paragraf dengan font Times New Roman dan aturan non-KBBI italic."""
    runs_data = parse_runs_for_italics(text)
    if not runs_data:
        runs_data = [(text, base_bold, base_italic)]
        
    for chunk, bold, italic in runs_data:
        run = paragraph.add_run(chunk)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(font_size_pt)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.bold = base_bold or bold
        run.italic = base_italic or italic

def setup_academic_numbering(doc):
    """
    Mendaftarkan definisi template penomoran berhierarki native Word (OpenXML) ke dalam numbering.xml:
    - Pola 1 (Utama, base 2.54 cm): ind left=1865, hanging=425 (angka di 2.54 cm, teks di 3.29 cm)
    - Pola 2 (Pengantar, base 1.27 cm): ind left=1145, hanging=425 (angka di 1.27 cm, teks di 2.02 cm)
    - Multilevel:
      - Level 0 (Level 1 user): 1. 2. 3. (decimal, %1.)
      - Level 1 (Level 2 user): a. b. c. (lowerLetter, %2.)
      - Level 2 (Level 3 user): 1) 2) 3) (decimal, %3))
    - Bullet list native:  (bullet), o (circle),  (square)
    """
    num_part = doc.part.numbering_part
    num_part_el = num_part._element

    if getattr(doc, '_academic_pola1_abs_id', None) is not None:
        return doc._academic_pola1_abs_id, doc._academic_pola2_abs_id, doc._academic_bullet_abstract_num_id

    max_abs_id = 0
    for abs_el in num_part_el.findall(qn('w:abstractNum')):
        try:
            val = int(abs_el.get(qn('w:abstractNumId')))
            if val > max_abs_id:
                max_abs_id = val
        except Exception:
            pass

    pola1_abs_id = max(100, max_abs_id + 1)
    pola2_abs_id = pola1_abs_id + 1
    bullet_abs_id = pola2_abs_id + 1

    # XML AbstractNum Pola 1 (1865 left, 425 hanging -> 2.54 cm bullet, 3.29 cm teks)
    pola1_abstract_xml = f'''
    <w:abstractNum xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:abstractNumId="{pola1_abs_id}">
      <w:multiLevelType w:val="hybridMultilevel"/>
      <w:tmpl w:val="A1B2C3D4"/>
      <w:lvl w:ilvl="0">
        <w:start w:val="1"/>
        <w:numFmt w:val="decimal"/>
        <w:lvlText w:val="%1."/>
        <w:lvlJc w:val="left"/>
        <w:pPr>
          <w:ind w:left="1865" w:hanging="425"/>
        </w:pPr>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
          <w:sz w:val="24"/>
          <w:color w:val="000000"/>
        </w:rPr>
      </w:lvl>
      <w:lvl w:ilvl="1">
        <w:start w:val="1"/>
        <w:numFmt w:val="lowerLetter"/>
        <w:lvlText w:val="%2."/>
        <w:lvlJc w:val="left"/>
        <w:pPr>
          <w:ind w:left="2290" w:hanging="425"/>
        </w:pPr>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
          <w:sz w:val="24"/>
          <w:color w:val="000000"/>
        </w:rPr>
      </w:lvl>
      <w:lvl w:ilvl="2">
        <w:start w:val="1"/>
        <w:numFmt w:val="decimal"/>
        <w:lvlText w:val="%3)"/>
        <w:lvlJc w:val="left"/>
        <w:pPr>
          <w:ind w:left="2715" w:hanging="425"/>
        </w:pPr>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
          <w:sz w:val="24"/>
          <w:color w:val="000000"/>
        </w:rPr>
      </w:lvl>
    </w:abstractNum>
    '''

    # XML AbstractNum Pola 2 (1145 left, 425 hanging -> 1.27 cm bullet, 2.02 cm teks)
    pola2_abstract_xml = f'''
    <w:abstractNum xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:abstractNumId="{pola2_abs_id}">
      <w:multiLevelType w:val="hybridMultilevel"/>
      <w:tmpl w:val="A1B2C3D5"/>
      <w:lvl w:ilvl="0">
        <w:start w:val="1"/>
        <w:numFmt w:val="decimal"/>
        <w:lvlText w:val="%1."/>
        <w:lvlJc w:val="left"/>
        <w:pPr>
          <w:ind w:left="1145" w:hanging="425"/>
        </w:pPr>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
          <w:sz w:val="24"/>
          <w:color w:val="000000"/>
        </w:rPr>
      </w:lvl>
      <w:lvl w:ilvl="1">
        <w:start w:val="1"/>
        <w:numFmt w:val="lowerLetter"/>
        <w:lvlText w:val="%2."/>
        <w:lvlJc w:val="left"/>
        <w:pPr>
          <w:ind w:left="1570" w:hanging="425"/>
        </w:pPr>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
          <w:sz w:val="24"/>
          <w:color w:val="000000"/>
        </w:rPr>
      </w:lvl>
      <w:lvl w:ilvl="2">
        <w:start w:val="1"/>
        <w:numFmt w:val="decimal"/>
        <w:lvlText w:val="%3)"/>
        <w:lvlJc w:val="left"/>
        <w:pPr>
          <w:ind w:left="1995" w:hanging="425"/>
        </w:pPr>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
          <w:sz w:val="24"/>
          <w:color w:val="000000"/>
        </w:rPr>
      </w:lvl>
    </w:abstractNum>
    '''

    # XML AbstractNum Bullet
    bullet_abstract_xml = f'''
    <w:abstractNum xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:abstractNumId="{bullet_abs_id}">
      <w:multiLevelType w:val="hybridMultilevel"/>
      <w:tmpl w:val="B1C2D3E4"/>
      <w:lvl w:ilvl="0">
        <w:start w:val="1"/>
        <w:numFmt w:val="bullet"/>
        <w:lvlText w:val=""/>
        <w:lvlJc w:val="left"/>
        <w:pPr>
          <w:ind w:left="1865" w:hanging="425"/>
        </w:pPr>
        <w:rPr>
          <w:rFonts w:ascii="Symbol" w:hAnsi="Symbol" w:hint="default"/>
          <w:sz w:val="24"/>
          <w:color w:val="000000"/>
        </w:rPr>
      </w:lvl>
      <w:lvl w:ilvl="1">
        <w:start w:val="1"/>
        <w:numFmt w:val="bullet"/>
        <w:lvlText w:val="o"/>
        <w:lvlJc w:val="left"/>
        <w:pPr>
          <w:ind w:left="2290" w:hanging="425"/>
        </w:pPr>
        <w:rPr>
          <w:rFonts w:ascii="Courier New" w:hAnsi="Courier New" w:hint="default"/>
          <w:sz w:val="20"/>
          <w:color w:val="000000"/>
        </w:rPr>
      </w:lvl>
      <w:lvl w:ilvl="2">
        <w:start w:val="1"/>
        <w:numFmt w:val="bullet"/>
        <w:lvlText w:val=""/>
        <w:lvlJc w:val="left"/>
        <w:pPr>
          <w:ind w:left="2715" w:hanging="425"/>
        </w:pPr>
        <w:rPr>
          <w:rFonts w:ascii="Wingdings" w:hAnsi="Wingdings" w:hint="default"/>
          <w:sz w:val="24"/>
          <w:color w:val="000000"/>
        </w:rPr>
      </w:lvl>
    </w:abstractNum>
    '''

    first_num = num_part_el.find(qn('w:num'))
    pola1_elem = parse_xml(pola1_abstract_xml)
    pola2_elem = parse_xml(pola2_abstract_xml)
    bullet_elem = parse_xml(bullet_abstract_xml)
    if first_num is not None:
        first_num.addprevious(pola1_elem)
        first_num.addprevious(pola2_elem)
        first_num.addprevious(bullet_elem)
    else:
        num_part_el.append(pola1_elem)
        num_part_el.append(pola2_elem)
        num_part_el.append(bullet_elem)

    doc._academic_pola1_abs_id = pola1_abs_id
    doc._academic_pola2_abs_id = pola2_abs_id
    doc._academic_bullet_abstract_num_id = bullet_abs_id
    return pola1_abs_id, pola2_abs_id, bullet_abs_id

def create_numbered_list_id(doc, pattern: int = 1):
    """
    Membuat ID penomoran (w:numId) baru agar daftar angka/huruf mulai dari awal (1.).
    pattern=1: Pola 1 (base 2.54 cm -> 3.29 cm teks)
    pattern=2: Pola 2 (base 1.27 cm -> 2.02 cm teks)
    """
    pola1_abs_id, pola2_abs_id, _ = setup_academic_numbering(doc)
    abs_id = pola1_abs_id if pattern == 1 else pola2_abs_id
    num_part = doc.part.numbering_part
    num_part_el = num_part._element

    max_num_id = 0
    for num_el in num_part_el.findall(qn('w:num')):
        try:
            val = int(num_el.get(qn('w:numId')))
            if val > max_num_id:
                max_num_id = val
        except Exception:
            pass

    new_num_id = max(100, max_num_id + 1)
    num_xml = f'''
    <w:num xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:numId="{new_num_id}">
      <w:abstractNumId w:val="{abs_id}"/>
      <w:lvlOverride w:ilvl="0">
        <w:startOverride w:val="1"/>
      </w:lvlOverride>
    </w:num>
    '''
    num_part_el.append(parse_xml(num_xml))
    return new_num_id

def create_bullet_list_id(doc, pattern: int = 1):
    """
    Membuat ID list bullet (w:numId) baru untuk simbol peluru.
    """
    _, _, bullet_abs_id = setup_academic_numbering(doc)
    num_part = doc.part.numbering_part
    num_part_el = num_part._element

    max_num_id = 0
    for num_el in num_part_el.findall(qn('w:num')):
        try:
            val = int(num_el.get(qn('w:numId')))
            if val > max_num_id:
                max_num_id = val
        except Exception:
            pass

    new_num_id = max(100, max_num_id + 1)
    num_xml = f'''
    <w:num xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:numId="{new_num_id}">
      <w:abstractNumId w:val="{bullet_abs_id}"/>
    </w:num>
    '''
    num_part_el.append(parse_xml(num_xml))
    return new_num_id

def create_document(template_path=None):
    """Membuat dokumen baru dari template atau default dengan margin 4-4-3-3 cm A4."""
    if template_path and os.path.exists(template_path):
        doc = docx.Document(template_path)
        # Bersihkan sisa paragraf awal template
        for p in list(doc.paragraphs):
            p._element.getparent().remove(p._element)
    else:
        doc = docx.Document()
        
    # Set properti halaman A4 & margin akademik 4-4-3-3 cm
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(4.0)
    section.left_margin = Cm(4.0)
    section.bottom_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
    
    # Inisialisasi konfigurasi native numbering akademik
    setup_academic_numbering(doc)
    return doc

def add_identity_header(doc, module_title: str, nama: str, nim: str, kelas: str, standalone_cover: bool = True):
    """
    Menambahkan halaman sampul (Cover Page) mandiri di Halaman 1 sesuai format referensi:
    - Judul dan Identitas diposisikan di area tengah halaman secara vertikal
    - Perataan Tengah (Center Aligned)
    - Times New Roman 12 pt, Bold untuk judul modul
    """
    if standalone_cover:
        # Tambahkan jeda vertikal ke tengah halaman (seperti Cover Laporan Praktikum referensi)
        for _ in range(10):
            p_blank = doc.add_paragraph()
            set_paragraph_spacing(p_blank, line_spacing=1.0, space_before_pt=0, space_after_pt=0)

    # Baris Judul
    lines = [l.strip() for l in module_title.strip().split("\n") if l.strip()]
    for i, line in enumerate(lines):
        p_title = doc.add_paragraph()
        if standalone_cover:
            p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        after_pt = 6 if i == len(lines) - 1 else 0
        set_paragraph_spacing(p_title, line_spacing=1.5, space_before_pt=0, space_after_pt=after_pt)
        r_title = p_title.add_run(line.upper() if line.isupper() or not line.startswith("(") else line)
        r_title.font.name = 'Times New Roman'
        r_title.font.size = Pt(12)
        r_title.bold = True
        r_title.font.color.rgb = RGBColor(0, 0, 0)
    
    # Baris Identitas
    for label, val in [("Nama : ", nama), ("NIM : ", nim), ("Kelas : ", kelas)]:
        p = doc.add_paragraph()
        if standalone_cover:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, line_spacing=1.5, space_before_pt=0, space_after_pt=0)
        r = p.add_run(f"{label}{val}")
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(0, 0, 0)
        
    if not standalone_cover:
        # Beri sedikit jeda paragraf kosong
        p_sep = doc.add_paragraph()
        set_paragraph_spacing(p_sep, line_spacing=1.0, space_before_pt=0, space_after_pt=6)

def add_heading_1(doc, text: str):
    """Menambahkan Heading 1 (contoh: BAB V PENUTUP, DAFTAR PUSTAKA) - Center, Bold."""
    p = doc.add_paragraph()
    try:
        p.style = doc.styles['Heading 1']
    except Exception:
        pass
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, line_spacing=1.5, space_before_pt=12, space_after_pt=6, keep_with_next=True)
    set_outline_level(p, level=0)
    add_runs(p, text, is_heading=True, base_bold=True, font_size_pt=12)
    return p

def add_heading_2(doc, text: str):
    """Menambahkan Heading 2 (contoh: 5.1 Kesimpulan) - Left, Bold."""
    p = doc.add_paragraph()
    try:
        p.style = doc.styles['Heading 2']
    except Exception:
        pass
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, line_spacing=1.5, space_before_pt=12, space_after_pt=4, keep_with_next=True)
    set_outline_level(p, level=1)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(1.27))
    add_runs(p, text, is_heading=True, base_bold=True, font_size_pt=12)
    return p

def add_heading_3(doc, text: str):
    """Menambahkan Heading 3 (contoh: 2.1.5 Judul, 3.1.5 Judul, 5.1.1 Judul) - Left, Bold."""
    p = doc.add_paragraph()
    try:
        p.style = doc.styles['Heading 3']
    except Exception:
        pass
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, line_spacing=1.5, space_before_pt=10, space_after_pt=4, keep_with_next=True)
    set_outline_level(p, level=2)
    if text.startswith("\t"):
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.27))
        p.paragraph_format.tab_stops.add_tab_stop(Cm(2.54))
    else:
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.27))
    add_runs(p, text, is_heading=True, base_bold=True, font_size_pt=12)
    return p

def add_body_paragraph(doc, text: str, first_line_indent_cm=1.27, left_indent_cm=0.0):
    """Menambahkan paragraf isi dengan indentasi baris pertama dan rata kanan-kiri (Justified)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, line_spacing=1.5, space_before_pt=0, space_after_pt=4)
    if left_indent_cm > 0:
        p.paragraph_format.left_indent = Cm(left_indent_cm)
    if first_line_indent_cm != 0:
        p.paragraph_format.first_line_indent = Cm(first_line_indent_cm)
    add_runs(p, text)
    return p

def add_page_break(doc):
    """Menambahkan pemisah halaman (Page Break) agar bab baru berada di halaman tersendiri."""
    if doc.paragraphs:
        return doc.paragraphs[-1].add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
    return doc.add_page_break()

def add_list_item(doc, text: str, prefix: str = "", level: int = 1, num_id: int = None, list_type: str = "number", base_indent_cm: float = 2.54, text_indent_cm: float = 0.75):
    """
    Menambahkan item list native Microsoft Word:
    - Menyematkan OpenXML <w:numPr> agar tombol Numbering / Bullets di ribbon Home Word MENYALA.
    - Mendukung hierarki multilevel:
      - level 1: 1., 2., 3.
      - level 2: a., b., c.
      - level 3: 1), 2), 3)
    - Auto-continue di Word (tekan Enter otomatis lanjut ke nomor berikutnya).
    - Text indent hanging 0.75 cm, nomor sejajar base_indent_cm (default 2.54 cm).
    """
    # Bersihkan prefix manual jika ada di dalam text
    clean_text = text.strip()
    clean_text = re.sub(r'^(?:[0-9]+\.|\b[a-zA-Z]\.|\b[0-9]+\)|\b[a-zA-Z]\)|[\u2022\u25E6\u25AA\-\*])\s*', '', clean_text)
    
    # Deteksi level dari prefix jika prefix dikirimkan
    if prefix:
        p_str = prefix.strip()
        if re.match(r'^[a-zA-Z]\.', p_str):
            level = max(level, 2)
        elif re.match(r'^[0-9]+\)', p_str) or re.match(r'^[a-zA-Z]\)', p_str):
            level = max(level, 3)
        elif p_str in ['-', '•', '*', '']:
            list_type = "bullet"

    p = doc.add_paragraph()
    try:
        p.style = doc.styles['List Paragraph']
    except Exception:
        pass
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, line_spacing=1.5, space_before_pt=0, space_after_pt=3)
    
    # Indentasi Word
    pos = base_indent_cm + (level - 1) * text_indent_cm
    p.paragraph_format.left_indent = Cm(pos + text_indent_cm)
    p.paragraph_format.first_line_indent = Cm(-text_indent_cm)
    
    # Dapatkan atau alokasikan num_id jika belum ada
    if num_id is None:
        pattern = 1 if base_indent_cm >= 2.0 else 2
        if list_type == "bullet":
            if not hasattr(doc, '_current_bullet_num_id'):
                doc._current_bullet_num_id = create_bullet_list_id(doc, pattern=pattern)
            num_id = doc._current_bullet_num_id
        else:
            if not hasattr(doc, '_current_numbered_num_id'):
                doc._current_numbered_num_id = create_numbered_list_id(doc, pattern=pattern)
            num_id = doc._current_numbered_num_id
            
    # Sematkan native OpenXML <w:numPr>
    pPr = p._element.get_or_add_pPr()
    ilvl_val = max(0, min(level - 1, 8))
    numPr_xml = f'<w:numPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:ilvl w:val="{ilvl_val}"/><w:numId w:val="{num_id}"/></w:numPr>'
    pPr.append(parse_xml(numPr_xml))
    
    # Tambahkan teks isi dengan cetak miring istilah non-KBBI
    add_runs(p, clean_text)
    return p

def add_list_body_paragraph(doc, text: str, level: int = 1, base_indent_cm: float = 2.54, text_indent_cm: float = 0.75, first_line_indent_cm: float = 0.75):
    """
    Menambahkan paragraf isi di bawah list item (misalnya uraian materi di bawah 1. Frame).
    left_indent sejajar teks list item (pos + text_indent_cm).
    first_line_indent menjorok 0.75 cm pada baris pertama jika > 0.
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, line_spacing=1.5, space_before_pt=0, space_after_pt=4)
    pos = base_indent_cm + (level - 1) * text_indent_cm
    p.paragraph_format.left_indent = Cm(pos + text_indent_cm)
    if first_line_indent_cm > 0:
        p.paragraph_format.first_line_indent = Cm(first_line_indent_cm)
    add_runs(p, text)
    return p

def add_conclusion_paragraph(doc, text: str, base_indent_cm: float = 1.27, first_line_offset_cm: float = 1.27):
    """
    Menambahkan paragraf kesimpulan di bawah sub-judul 5.1.N:
    - Baris pertama (kata pertama) sejajar dengan huruf pertama judul sub-bab (base_indent_cm + first_line_offset_cm = 2.54 cm).
    - Baris kedua dan seterusnya membungkus lurus sejajar dengan angka 5 (base_indent_cm = 1.27 cm).
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, line_spacing=1.5, space_before_pt=0, space_after_pt=4)
    p.paragraph_format.left_indent = Cm(base_indent_cm)
    p.paragraph_format.first_line_indent = Cm(first_line_offset_cm)
    add_runs(p, text)
    return p

def add_image_placeholder(doc, description: str, left_indent_cm: float = 0.0):
    """Menambahkan penanda placeholder untuk foto/gambar yang akan disisipkan user."""
    p = doc.add_paragraph()
    if left_indent_cm > 0:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(left_indent_cm)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, line_spacing=1.5, space_before_pt=6, space_after_pt=6)
    placeholder_text = f"[Sisipkan Gambar: {description.strip()}]"
    run = p.add_run(placeholder_text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(100, 100, 100)
    run.italic = True
    run.bold = True
    return p

def add_bibliography_entry(doc, entry_text: str):
    """Menambahkan entri Daftar Pustaka dengan format hanging indent APA 7th Edition."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, line_spacing=1.5, space_before_pt=0, space_after_pt=6)
    # Hanging indent: margin kiri 1.27 cm, first line indent -1.27 cm
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.first_line_indent = Cm(-1.27)
    
    # Format APA: Hanya judul buku / jurnal dalam *...* yang dicetak miring, nama penulis & penerbit tetap reguler
    pattern = re.compile(r'(\*.*?\*)')
    parts = pattern.split(entry_text)
    for part in parts:
        if not part:
            continue
        is_it = part.startswith('*') and part.endswith('*') and len(part) >= 2
        chunk = part[1:-1] if is_it else part
        run = p.add_run(chunk)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.italic = is_it
    return p
